"""Streaming, read-only raw corpus EDA helpers for supporting notebooks."""

from __future__ import annotations

import hashlib
import json
import math
import re
import warnings
import zipfile
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterator

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook

SEED = 20260816
SAMPLE_MODULUS = 1000
SAMPLE_THRESHOLD = 10
HTML_RE = re.compile(r"<[A-Za-z/!][^>\n]{0,200}>")
CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
ZERO_WIDTH_RE = re.compile(r"[\u200b\u200c\u200d\u2060\ufeff]")
HANGUL_RE = re.compile(r"[가-힣ㄱ-ㅎㅏ-ㅣ]")
LATIN_RE = re.compile(r"[A-Za-z]")
COLORS = {"KO": "#2563EB", "EN": "#F97316", "neutral": "#64748B", "risk": "#DC2626"}


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def configure_plotting() -> None:
    plt.rcParams.update({
        "font.family": "NanumGothic",
        "axes.unicode_minus": False,
        "figure.facecolor": "#F8FAFC",
        "axes.facecolor": "white",
        "axes.titleweight": "bold",
        "axes.titlepad": 10,
        "axes.grid": True,
        "grid.alpha": 0.18,
        "savefig.dpi": 150,
    })


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_hash(parts: list[str], digest_size: int = 8) -> int:
    digest = hashlib.blake2b(digest_size=digest_size)
    digest.update(str(SEED).encode())
    for part in parts:
        encoded = part.encode("utf-8", errors="strict")
        digest.update(len(encoded).to_bytes(8, "big"))
        digest.update(encoded)
    return int.from_bytes(digest.digest(), "big")


def iter_json_data(path: Path) -> Iterator[dict[str, Any]]:
    decoder = json.JSONDecoder()
    with path.open("r", encoding="utf-8-sig", errors="strict") as handle:
        prefix = ""
        match = None
        while match is None:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                raise ValueError(f"data[] not found: {path}")
            prefix += chunk
            match = re.search(r'"data"\s*:\s*\[', prefix)
        buffer, position = prefix[match.end():], 0
        while True:
            while True:
                while position < len(buffer) and buffer[position] in " \t\r\n,":
                    position += 1
                if position < len(buffer):
                    break
                buffer, position = handle.read(4 * 1024 * 1024), 0
                if not buffer:
                    raise ValueError(f"unexpected EOF: {path}")
            if buffer[position] == "]":
                return
            while True:
                try:
                    value, end = decoder.raw_decode(buffer, position)
                    break
                except json.JSONDecodeError:
                    remainder = buffer[position:]
                    chunk = handle.read(4 * 1024 * 1024)
                    if not chunk:
                        raise
                    buffer, position = remainder + chunk, 0
            if not isinstance(value, dict):
                raise TypeError("data[] member is not an object")
            yield value
            position = end
            if position > 8 * 1024 * 1024:
                buffer, position = buffer[position:], 0


def value_type(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, str):
        return "string"
    if isinstance(value, (datetime, date)):
        return "datetime"
    return type(value).__name__


def label(value: Any) -> str:
    if value is None:
        return "<NULL>"
    if value == "":
        return "<EMPTY>"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def json_path_meta(relative: str) -> tuple[str, str, str]:
    split = "TRAIN" if "/1.Training/" in f"/{relative}" else "VALID" if "/2.Validation/" in f"/{relative}" else "UNSPLIT"
    role = "LABEL" if "/라벨링데이터/" in f"/{relative}" else "SOURCE" if "/원천데이터/" in f"/{relative}" else "OTHER"
    name = Path(relative).name
    direction = "EN_TO_KO" if "영한" in name else "KO_TO_EN" if "한영" in name else "UNKNOWN"
    return split, role, direction


def xlsx_family(path: Path) -> str:
    name = path.name
    if name.startswith("1_구어체"):
        return "구어체"
    if name.startswith("2_대화체"):
        return "대화체"
    if name.startswith("3_문어체_뉴스"):
        return "뉴스"
    if name.startswith("4_문어체_한국문화"):
        return "한국문화"
    if name.startswith("5_문어체_조례"):
        return "조례"
    if name.startswith("6_문어체_지자체"):
        return "지자체웹"
    return path.stem


def summarize_hist(counter: Counter[int]) -> dict[str, float | int]:
    total = sum(counter.values())
    if not total:
        return {"count": 0, "min": 0, "p50": 0, "p90": 0, "p95": 0, "p99": 0, "max": 0, "mean": 0.0}
    ordered = sorted(counter.items())
    def q(probability: float) -> int:
        target, cumulative = max(1, math.ceil(total * probability)), 0
        for value, count in ordered:
            cumulative += count
            if cumulative >= target:
                return value
        return ordered[-1][0]
    return {
        "count": total, "min": ordered[0][0], "p50": q(.5), "p90": q(.9),
        "p95": q(.95), "p99": q(.99), "max": ordered[-1][0],
        "mean": sum(value * count for value, count in ordered) / total,
    }


class CorpusAccumulator:
    def __init__(self) -> None:
        self.group_counts: Counter[str] = Counter()
        self.length_hist: dict[tuple[str, str, str], Counter[int]] = {}
        self.noise: dict[tuple[str, str], Counter[str]] = {}
        self.categories: Counter[tuple[str, str, str]] = Counter()
        self.schema: dict[tuple[str, str], dict[str, Any]] = {}
        self.samples: list[dict[str, Any]] = []
        self.seen_pairs: set[int] = set()
        self.duplicate_pair_rows = 0
        self.duplicate_pair_hashes: set[int] = set()
        self.training_pairs: set[int] = set()
        self.validation_pair_overlap = 0
        self.file_key_duplicates: Counter[str] = Counter()

    def add(self, row: dict[str, Any], *, group: str, relative: str, row_number: int,
            ko_key: str, en_key: str, pair_key: str | None, split: str,
            category_keys: list[str], file_seen_keys: set[int]) -> None:
        prior = self.group_counts[group]
        self.group_counts[group] += 1
        for field in row:
            key = (group, field)
            if key not in self.schema:
                self.schema[key] = {"absent": prior, "null": 0, "empty": 0, "types": Counter()}
        for (schema_group, field), state in self.schema.items():
            if schema_group != group:
                continue
            if field not in row:
                state["absent"] += 1
            else:
                value = row[field]
                state["types"][value_type(value)] += 1
                if value is None:
                    state["null"] += 1
                elif isinstance(value, str) and value == "":
                    state["empty"] += 1

        ko, en = str(row.get(ko_key) or ""), str(row.get(en_key) or "")
        for language, text in (("KO", ko), ("EN", en)):
            self.length_hist.setdefault((group, language, "codepoints"), Counter())[len(text)] += 1
            self.length_hist.setdefault((group, language, "utf8_bytes"), Counter())[len(text.encode("utf-8"))] += 1
            flags = self.noise.setdefault((group, language), Counter())
            tests = {
                "empty": text == "", "leading_or_trailing_ws": text != text.strip(),
                "repeated_space": "  " in text, "tab": "\t" in text, "newline": "\n" in text or "\r" in text,
                "html_like": bool(HTML_RE.search(text)), "control_char": bool(CONTROL_RE.search(text)),
                "zero_width": bool(ZERO_WIDTH_RE.search(text)), "replacement_char": "\ufffd" in text,
                "cross_script": bool(LATIN_RE.search(text)) if language == "KO" else bool(HANGUL_RE.search(text)),
                "expected_script_absent": bool(text) and (not HANGUL_RE.search(text) if language == "KO" else not LATIN_RE.search(text)),
            }
            for metric, present in tests.items():
                if present:
                    flags[metric] += 1

        pair_hash = stable_hash([ko, en])
        if pair_hash in self.seen_pairs:
            self.duplicate_pair_rows += 1
            self.duplicate_pair_hashes.add(pair_hash)
        else:
            self.seen_pairs.add(pair_hash)
        if split == "VALID" and pair_hash in self.training_pairs:
            self.validation_pair_overlap += 1
        elif split == "TRAIN":
            self.training_pairs.add(pair_hash)

        key_value = row.get(pair_key) if pair_key else None
        if key_value not in (None, ""):
            key_hash = stable_hash([label(key_value)])
            if key_hash in file_seen_keys:
                self.file_key_duplicates[relative] += 1
            else:
                file_seen_keys.add(key_hash)

        for category in category_keys:
            if category in row:
                self.categories[(group, category, label(row[category]))] += 1

        sample_key = label(key_value) if key_value not in (None, "") else str(row_number)
        if stable_hash([relative, sample_key]) % SAMPLE_MODULUS < SAMPLE_THRESHOLD:
            self.samples.append({
                "relative_path": relative, "group": group, "row_number": row_number,
                "pair_key": sample_key, "ko_preview": ko[:160], "en_preview": en[:160],
                "ko_codepoints": len(ko), "en_codepoints": len(en),
                "ko_utf8_bytes": len(ko.encode("utf-8")), "en_utf8_bytes": len(en.encode("utf-8")),
                **{key: label(row.get(key)) for key in category_keys[:5] if key in row},
            })


def iter_xlsx_rows(path: Path) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    raw_header = next(iterator, ())
    last = max((i for i, value in enumerate(raw_header) if value not in (None, "")), default=-1)
    header = [str(value) if value not in (None, "") else f"__unnamed_{i+1}" for i, value in enumerate(raw_header[:last + 1])]
    def rows() -> Iterator[tuple[int, dict[str, Any]]]:
        try:
            for row_number, values in enumerate(iterator, start=2):
                values = values[:len(header)]
                if not any(value not in (None, "") for value in values):
                    continue
                yield row_number, {header[i]: value for i, value in enumerate(values)}
        finally:
            workbook.close()
    return header, rows()


def profile_dataset(config: dict[str, Any]) -> dict[str, Any]:
    started = now_iso()
    raw_dir = Path(config["raw_dir"])
    paths = sorted((path for path in raw_dir.rglob("*") if path.is_file()), key=lambda p: p.relative_to(raw_dir).as_posix().encode())
    initial = {path: (path.stat().st_size, path.stat().st_mtime_ns) for path in paths}
    hashes = {path: sha256_file(path) for path in paths}
    first_by_hash: dict[str, Path] = {}
    inventory_rows = []
    for path in paths:
        rel = path.relative_to(raw_dir).as_posix()
        digest = hashes[path]
        duplicate_of = first_by_hash.get(digest)
        if duplicate_of is None:
            first_by_hash[digest] = path
        split, role, direction = json_path_meta(rel) if path.suffix.lower() == ".json" else ("UNSPLIT", "WORKBOOK" if path.suffix.lower() == ".xlsx" else "ARCHIVE", "KO_FIELD_TO_EN_FIELD" if path.suffix.lower() == ".xlsx" else "N/A")
        inventory_rows.append({
            "relative_path": rel, "bytes": initial[path][0], "mtime_epoch_ns": initial[path][1],
            "sha256": digest, "format": path.suffix.lower().lstrip(".").upper(),
            "split": split, "role": role, "direction": direction,
            "duplicate_of": duplicate_of.relative_to(raw_dir).as_posix() if duplicate_of else "",
        })
    inventory = pd.DataFrame(inventory_rows)
    unique_data_paths = [first_by_hash[digest] for digest in first_by_hash if first_by_hash[digest].suffix.lower() in {".json", ".xlsx"}]
    unique_data_paths.sort(key=lambda p: (0 if "/1.Training/" in f"/{p.relative_to(raw_dir).as_posix()}" else 1, p.relative_to(raw_dir).as_posix().encode()))
    acc = CorpusAccumulator()
    for path in unique_data_paths:
        relative = path.relative_to(raw_dir).as_posix()
        file_seen_keys: set[int] = set()
        if path.suffix.lower() == ".json":
            split, _, direction = json_path_meta(relative)
            group = f"{split} · {direction}"
            categories = ["data_set", "domain", "subdomain", "source", "style", "source_language", "target_language", "license"]
            for row_number, row in enumerate(iter_json_data(path), start=1):
                acc.add(row, group=group, relative=relative, row_number=row_number, ko_key="ko", en_key="en", pair_key="sn", split=split, category_keys=categories, file_seen_keys=file_seen_keys)
        else:
            header, rows = iter_xlsx_rows(path)
            group = xlsx_family(path)
            pair_key = "SID" if "SID" in header else "ID" if "ID" in header else None
            categories = [field for field in header if field not in {"SID", "ID", "원문", "번역문", "URL", "날짜", "Set Nr."}]
            for row_number, row in rows:
                acc.add(row, group=group, relative=relative, row_number=row_number, ko_key="원문", en_key="번역문", pair_key=pair_key, split="UNSPLIT", category_keys=categories, file_seen_keys=file_seen_keys)

    final = {path: (path.stat().st_size, path.stat().st_mtime_ns) for path in paths}
    changed = [path.relative_to(raw_dir).as_posix() for path in paths if initial[path] != final[path]]
    if changed:
        raise RuntimeError(f"DOWNLOAD_IN_PROGRESS, excluded snapshot: {changed}")
    inventory["stability_state"] = "STABLE_FILE"

    schema_rows = []
    for (group, field), state in sorted(acc.schema.items()):
        total = acc.group_counts[group]
        schema_rows.append({
            "group": group, "field": field, "records": total, "absent": state["absent"],
            "null": state["null"], "empty": state["empty"],
            "missing_or_empty_rate": (state["absent"] + state["null"] + state["empty"]) / total if total else 0,
            "type_counts": json.dumps(dict(state["types"]), ensure_ascii=False, sort_keys=True),
        })
    length_rows = []
    for (group, language, metric), counter in sorted(acc.length_hist.items()):
        length_rows.append({"group": group, "language": language, "metric": metric, **summarize_hist(counter)})
    noise_rows = []
    for (group, language), counter in sorted(acc.noise.items()):
        total = acc.group_counts[group]
        for metric, count in sorted(counter.items()):
            noise_rows.append({"group": group, "language": language, "metric": metric, "row_count": count, "records": total, "rows_per_100k": count / total * 100_000})
    category_rows = [{"group": group, "dimension": dimension, "value": value, "count": count} for (group, dimension, value), count in acc.categories.items()]
    samples = pd.DataFrame(acc.samples)
    duplicate_rows = [{
        "scope": "dataset", "duplicate_pair_rows_after_first": acc.duplicate_pair_rows,
        "distinct_duplicate_pair_hashes": len(acc.duplicate_pair_hashes),
        "validation_rows_overlapping_training_pair": acc.validation_pair_overlap,
        "candidate_method": "BLAKE2b-64 raw KO+EN; collision possible; no normalization",
    }]
    for relative, count in acc.file_key_duplicates.items():
        duplicate_rows.append({"scope": relative, "duplicate_pair_rows_after_first": np.nan, "distinct_duplicate_pair_hashes": np.nan, "validation_rows_overlapping_training_pair": np.nan, "duplicate_nonempty_pair_key_rows": count, "candidate_method": "file-scoped raw key"})

    summary = {
        "dataset_local_id": config["dataset_local_id"], "dataset_title": config["dataset_title"],
        "audit_timestamp": started, "audit_completed_timestamp": now_iso(),
        "raw_dir": str(raw_dir), "physical_file_count": len(paths),
        "unique_content_file_count": len(first_by_hash), "profiled_unique_data_files": len(unique_data_paths),
        "logical_record_count": int(sum(acc.group_counts.values())), "sample_record_count": len(samples),
        "population_scope": "FULL POPULATION AGGREGATES over unique SHA-256 JSON/XLSX content",
        "sample_method": f"deterministic BLAKE2b(seed={SEED}) hash threshold {SAMPLE_THRESHOLD}/{SAMPLE_MODULUS} (~1%); scatter/sample only",
        "notebook_status": "SUPPORTING / EXPLORATORY / PRE-INGESTION EDA / NOT G1 PASS",
        "raw_file_manifest_sha": hashlib.sha256("".join(f"{row['relative_path']}\t{row['bytes']}\t{row['mtime_epoch_ns']}\t{row['sha256']}\n" for row in inventory_rows).encode()).hexdigest(),
        "stability_counts": {"STABLE_FILE": len(paths), "DOWNLOAD_IN_PROGRESS": 0, "PARTIAL_DOWNLOAD": 0, "UNKNOWN": 0},
    }
    return {
        "config": config, "summary": summary, "inventory": inventory,
        "record_counts": pd.DataFrame([{"group": group, "records": count} for group, count in acc.group_counts.items()]),
        "schema": pd.DataFrame(schema_rows), "length_summary": pd.DataFrame(length_rows),
        "noise": pd.DataFrame(noise_rows), "categories": pd.DataFrame(category_rows),
        "duplicates": pd.DataFrame(duplicate_rows), "samples": samples,
        "length_hist": acc.length_hist,
    }


def save_artifacts(result: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix = result["config"]["artifact_prefix"]
    for key in ["inventory", "record_counts", "schema", "length_summary", "noise", "categories", "duplicates", "samples"]:
        result[key].to_csv(output_dir / f"{prefix}_{key}.csv", index=False, encoding="utf-8-sig")
    (output_dir / f"{prefix}_profile_summary.json").write_text(json.dumps(result["summary"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _save(fig: plt.Figure, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, bbox_inches="tight")
    plt.show()


def plot_overview(result: dict[str, Any], figure_dir: Path) -> None:
    configure_plotting()
    prefix, title = result["config"]["artifact_prefix"], result["config"]["dataset_title"]
    fig, axes = plt.subplots(2, 2, figsize=(15, 10), constrained_layout=True)
    counts = result["record_counts"].sort_values("records")
    axes[0, 0].barh(counts["group"], counts["records"], color="#0F766E")
    axes[0, 0].xaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{x/1e6:.1f}M" if x >= 1e6 else f"{x/1e3:.0f}k"))
    axes[0, 0].set_title("Full population: logical record scale")
    axes[0, 0].set_xlabel("translation-pair records")
    inv = result["inventory"].copy()
    inv["MiB"] = inv["bytes"] / 2**20
    inv["file"] = inv["relative_path"].map(lambda x: Path(x).name[:34])
    axes[0, 1].barh(inv["file"], inv["MiB"], color=np.where(inv["duplicate_of"].eq(""), "#334155", "#CBD5E1"))
    axes[0, 1].set_title("Physical files (light = byte-identical alias)")
    axes[0, 1].set_xlabel("MiB")
    schema = result["schema"]
    key_fields = [field for field in ["ko", "en", "sn", "원문", "번역문", "SID", "ID", "domain", "source", "ner", "URL", "자동분류2", "자동분류3"] if field in set(schema["field"])]
    matrix = schema[schema["field"].isin(key_fields)].pivot(index="field", columns="group", values="missing_or_empty_rate").fillna(1.0)
    image = axes[1, 0].imshow(matrix.values, aspect="auto", cmap="YlOrRd", vmin=0, vmax=max(.01, float(matrix.values.max())))
    axes[1, 0].set_yticks(range(len(matrix.index)), matrix.index)
    axes[1, 0].set_xticks(range(len(matrix.columns)), matrix.columns, rotation=25, ha="right")
    axes[1, 0].set_title("Field missing/null/empty rate (white ≈ complete)")
    fig.colorbar(image, ax=axes[1, 0], shrink=.75, label="rate")
    summary = result["summary"]
    axes[1, 1].axis("off")
    axes[1, 1].text(0, .95, "Snapshot contract", fontsize=16, fontweight="bold", va="top")
    axes[1, 1].text(0, .82, "\n".join([
        f"Local ID: {summary['dataset_local_id']}", f"Physical files: {summary['physical_file_count']}",
        f"Unique content files: {summary['unique_content_file_count']}", f"Logical records: {summary['logical_record_count']:,}",
        f"Deterministic sample: {summary['sample_record_count']:,}", f"Manifest SHA: {summary['raw_file_manifest_sha'][:16]}…",
        "Status: SUPPORTING / EXPLORATORY", "Claim boundary: NOT G1 PASS",
    ]), fontsize=12, va="top", linespacing=1.55)
    fig.suptitle(f"{title}\nStructure, scale, and traceability", fontsize=18, fontweight="bold")
    _save(fig, figure_dir / f"{prefix}_01_structure_scale.png")


def plot_lengths(result: dict[str, Any], figure_dir: Path) -> None:
    configure_plotting()
    prefix, title = result["config"]["artifact_prefix"], result["config"]["dataset_title"]
    fig, axes = plt.subplots(2, 2, figsize=(15, 10), constrained_layout=True)
    groups = list(result["record_counts"].sort_values("group")["group"])
    for axis, language in zip(axes[0], ["KO", "EN"]):
        for group in groups:
            counter = result["length_hist"].get((group, language, "codepoints"), Counter())
            summary = summarize_hist(counter)
            limit = max(1, int(summary["p99"]))
            xs = np.arange(limit + 1)
            ys = np.array([counter.get(int(x), 0) for x in xs], dtype=float)
            if ys.sum():
                axis.plot(xs, ys / ys.sum(), label=group, linewidth=1.7)
        axis.set_title(f"{language} raw codepoint length (full population, ≤ group p99)")
        axis.set_xlabel("Unicode codepoints")
        axis.set_ylabel("within-group share")
        axis.legend(fontsize=8, loc="upper right")
    summary = result["length_summary"]
    codepoints = summary[summary["metric"].eq("codepoints")].copy()
    codepoints["label"] = codepoints["group"] + " · " + codepoints["language"]
    y = np.arange(len(codepoints))
    axes[1, 0].barh(y, codepoints["p95"], color=[COLORS.get(lang, COLORS["neutral"]) for lang in codepoints["language"]], alpha=.35, label="p95")
    axes[1, 0].scatter(codepoints["p50"], y, color="#111827", s=30, label="p50")
    axes[1, 0].set_yticks(y, codepoints["label"])
    axes[1, 0].set_xlabel("raw codepoints")
    axes[1, 0].set_title("Median dot within p95 bar")
    axes[1, 0].legend()
    samples = result["samples"]
    if len(samples):
        hb = axes[1, 1].hexbin(samples["ko_codepoints"], samples["en_codepoints"], gridsize=42, mincnt=1, cmap="viridis")
        fig.colorbar(hb, ax=axes[1, 1], label="sample rows per hex")
    axes[1, 1].set_xlabel("KO raw codepoints")
    axes[1, 1].set_ylabel("EN raw codepoints")
    axes[1, 1].set_title(f"Deterministic ~1% sample relationship (n={len(samples):,})")
    fig.suptitle(f"{title}\nRaw KO/EN length structure", fontsize=18, fontweight="bold")
    _save(fig, figure_dir / f"{prefix}_02_raw_lengths.png")


def plot_categories(result: dict[str, Any], figure_dir: Path) -> None:
    configure_plotting()
    prefix, title = result["config"]["artifact_prefix"], result["config"]["dataset_title"]
    categories = result["categories"]
    available = set(categories["dimension"])
    preferred = result["config"].get("category_plot_dimensions", [])
    dimensions = [dimension for dimension in preferred if dimension in available][:4]
    if len(dimensions) < 4:
        fallback = categories.groupby("dimension")["count"].sum().sort_values(ascending=False).index.tolist()
        dimensions.extend(dimension for dimension in fallback if dimension not in dimensions)
        dimensions = dimensions[:4]
    fig, axes = plt.subplots(2, 2, figsize=(15, 10), constrained_layout=True)
    for axis, dimension in zip(axes.flat, dimensions):
        top = categories[categories["dimension"].eq(dimension)].groupby("value")["count"].sum().nlargest(10).sort_values()
        axis.barh([str(value)[:38] for value in top.index], top.values, color="#7C3AED")
        axis.set_title(f"{dimension}: top raw categories")
        axis.set_xlabel("records (full population)")
    for axis in list(axes.flat)[len(dimensions):]:
        axis.axis("off")
    fig.suptitle(f"{title}\nDomain, source, and category composition", fontsize=18, fontweight="bold")
    _save(fig, figure_dir / f"{prefix}_03_category_composition.png")


def plot_quality(result: dict[str, Any], figure_dir: Path) -> None:
    configure_plotting()
    prefix, title = result["config"]["artifact_prefix"], result["config"]["dataset_title"]
    fig, axes = plt.subplots(2, 2, figsize=(15, 10), constrained_layout=True)
    noise = result["noise"].groupby(["language", "metric"], as_index=False)[["row_count", "records"]].sum()
    noise["rows_per_100k"] = noise["row_count"] / noise["records"] * 100_000
    top_noise = noise.nlargest(14, "rows_per_100k").sort_values("rows_per_100k")
    axes[0, 0].barh(top_noise["language"] + " · " + top_noise["metric"], top_noise["rows_per_100k"], color=[COLORS.get(x, "#64748B") for x in top_noise["language"]])
    axes[0, 0].set_title("Raw noise/script signals (not automatic errors)")
    axes[0, 0].set_xlabel("flagged rows per 100k")
    schema = result["schema"].copy()
    schema["missing_count"] = schema["absent"] + schema["null"] + schema["empty"]
    top_missing = schema[schema["missing_count"].gt(0)].nlargest(12, "missing_or_empty_rate").sort_values("missing_or_empty_rate")
    axes[0, 1].barh(top_missing["group"] + " · " + top_missing["field"], top_missing["missing_or_empty_rate"], color="#D97706")
    axes[0, 1].set_xlim(0, 1)
    axes[0, 1].set_title("Highest metadata missing/null/empty rates")
    axes[0, 1].set_xlabel("rate")
    duplicates = result["duplicates"].iloc[0]
    dup_values = [duplicates.get("duplicate_pair_rows_after_first", 0), duplicates.get("validation_rows_overlapping_training_pair", 0)]
    axes[1, 0].bar(["pair duplicate\ncandidates", "validation→training\noverlap candidates"], dup_values, color=["#DC2626", "#EA580C"])
    axes[1, 0].set_title("Full-population raw hash signals")
    axes[1, 0].set_ylabel("rows after first / overlap rows")
    axes[1, 1].axis("off")
    axes[1, 1].text(0, .95, "Interpretation boundary", fontsize=16, fontweight="bold", va="top")
    axes[1, 1].text(0, .80, "\n".join([
        "• BLAKE2b-64 counts are duplicate candidates.", "• Script mixing may be legitimate names or quotations.",
        "• HTML-like patterns are regex observations.", "• Missing metadata does not invalidate a text pair.",
        "• All signals require reconciliation before QC decisions.", "• No normalization was persisted.",
    ]), fontsize=12, va="top", linespacing=1.6)
    fig.suptitle(f"{title}\nPotential QC signals and metadata coverage", fontsize=18, fontweight="bold")
    _save(fig, figure_dir / f"{prefix}_04_quality_signals.png")


def run_and_render(config: dict[str, Any], output_dir: Path, figure_dir: Path) -> dict[str, Any]:
    result = profile_dataset(config)
    save_artifacts(result, output_dir)
    plot_overview(result, figure_dir)
    plot_lengths(result, figure_dir)
    plot_categories(result, figure_dir)
    plot_quality(result, figure_dir)
    return result
